"""Abertura do template Claro e criação de slides em branco."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from lib import theme


def open_template(path: Path = theme.TEMPLATE_PATH) -> Presentation:
    """Abre o template e remove todos os slides exceto a capa."""
    prs = Presentation(str(path))
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids)[1:]:
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)
    return prs


def new_slide(prs: Presentation):
    """Adiciona slide com o layout em branco do template."""
    return prs.slides.add_slide(prs.slide_layouts[theme.TEMPLATE_BLANK_LAYOUT])


def replace_cover_text(prs: Presentation, old: str, new: str) -> None:
    """Substitui um trecho de texto nos runs da capa."""
    for shape in prs.slides[0].shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
