"""Primitivas de desenho (caixas, textos, cabeçalho, rodapé e capturas) sobre python-pptx."""

from __future__ import annotations

from pathlib import Path
from typing import Union

from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from lib import theme

Run = tuple[str, dict]
Paragraph = Union[str, list[Run]]
Paragraphs = Union[str, list[Paragraph]]


def rgb(hex_color: str) -> RGBColor:
    """Converte cor hexadecimal (sem #) em RGBColor."""
    return RGBColor.from_string(hex_color)


def add_box(slide, x: int, y: int, w: int, h: int, fill: str | None = None, line: str | None = None,
            shape=MSO_SHAPE.RECTANGLE, line_width: float = 0.75):
    """Adiciona forma com preenchimento e borda opcionais, sem sombra."""
    shp = slide.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_width)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = 0.12
    return shp


def add_text(slide, x: int, y: int, w: int, h: int, paras: Paragraphs, size: float = 12,
             color: str = theme.BODY, bold: bool = False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             spacing: float = 4, shape=None, margin: int = 45720):
    """Escreve parágrafos numa caixa nova ou em `shape`; cada parágrafo é str ou lista de (texto, formato)."""
    if shape is None:
        shape = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Emu(margin)
    frame.margin_top = frame.margin_bottom = Emu(27432)
    frame.vertical_anchor = anchor
    for index, para in enumerate([paras] if isinstance(paras, str) else paras):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(spacing)
        for content, fmt in (para if isinstance(para, list) else [(para, {})]):
            run = paragraph.add_run()
            run.text = content
            run.font.name = theme.FONT
            run.font.size = Pt(fmt.get("size", size))
            run.font.bold = fmt.get("bold", bold)
            run.font.color.rgb = rgb(fmt.get("color", color))
    return shape


def add_label(shape, content: Paragraphs, size: float = 12, color: str = theme.NAVY, bold: bool = True,
              align=PP_ALIGN.LEFT, margin: int = 45720, spacing: float = 0):
    """Escreve texto centralizado verticalmente dentro de uma forma existente."""
    return add_text(None, 0, 0, 0, 0, content, size=size, color=color, bold=bold, align=align,
                    anchor=MSO_ANCHOR.MIDDLE, shape=shape, margin=margin, spacing=spacing)


def bullet_paragraphs(items: list[str], size: float = 12.5, color: str = theme.BODY) -> list[Paragraph]:
    """Monta parágrafos com marcador vermelho para `add_text`."""
    return [[("•  ", {"color": theme.RED, "bold": True, "size": size}), (item, {"size": size, "color": color})]
            for item in items]


def add_header(slide, tag: str, title: str, subtitle: str | None = None) -> None:
    """Desenha etiqueta da seção, título, subtítulo opcional e linha divisória."""
    add_text(slide, theme.MARGIN_X, 338328, 6000000, 260000, tag, size=10, color=theme.RED, bold=True)
    add_text(slide, theme.MARGIN_X, 620000, 10000000, 560000, title, size=30, color=theme.NAVY, bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_text(slide, theme.MARGIN_X, 1190000, 10300000, 430000, subtitle, size=13)
    add_box(slide, theme.MARGIN_X, 1690000, theme.CONTENT_WIDTH, 18288, fill=theme.LINE)


def add_logo(slide) -> None:
    """Insere o logo Claro no canto inferior direito."""
    x, y, w, h = theme.LOGO_BOX
    slide.shapes.add_picture(str(theme.LOGO_PATH), Emu(x), Emu(y), Emu(w), Emu(h))


def add_footer(slide, message: str, logo: bool = True) -> None:
    """Desenha a faixa vermelha com a mensagem-chave do slide e o logo."""
    add_box(slide, theme.MARGIN_X, theme.FOOTER_TOP, theme.CONTENT_WIDTH, 32004, fill=theme.RED)
    add_text(slide, theme.MARGIN_X, 5880000, 10100000, 380000, message, size=13, color=theme.NAVY, bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
    if logo:
        add_logo(slide)


def set_notes(slide, content: str) -> None:
    """Grava o roteiro do apresentador nas notas do slide."""
    slide.notes_slide.notes_text_frame.text = content


def add_screenshot(slide, path: Path, x: int, y: int, w: int | None = None, h: int | None = None,
                   label: str | None = None) -> tuple[int, int]:
    """Insere captura com borda, preservando a proporção do PNG; retorna (largura, altura) em EMU."""
    with Image.open(path) as image:
        ratio = image.size[0] / image.size[1]
    w = w or int(h * ratio)
    h = h or int(w / ratio)
    picture = slide.shapes.add_picture(str(path), Emu(x), Emu(y), Emu(w), Emu(h))
    picture.line.color.rgb = rgb(theme.LINE)
    picture.line.width = Pt(0.75)
    if label:
        tag = add_box(slide, x + 91440, y + h - 391440, len(label) * 70000 + 228600, 300000, fill=theme.NAVY)
        add_label(tag, label, size=10.5, color=theme.WHITE, margin=91440)
    return w, h
